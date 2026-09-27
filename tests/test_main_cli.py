import csv
import sys
from pathlib import Path

import pytest
import yaml

from src.main import main

PROJECT_ROOT = Path(__file__).resolve().parents[1]

HEADERS = [
    "采购合同号", "客户名称", "玉湖主体", "供应商名称", "序号", "合同产品名称",
    "国家", "厂号", "柜号", "规格", "件数/箱", "重量", "单价/kg", "金额/元",
    "生产日期", "仓库（货物位置）",
]
ROW = [
    "YH2026P001", "客户A", "玉湖主体", "供应商B", "1", "牛肉",
    "巴西", "SIF123", "C1", "20kg/箱", "100", "2000", "50", "100000",
    "2025-10-15", "广州仓",
]



@pytest.fixture
def excel_file(tmp_path):
    """第 0 行标题、第 1 行表头、第 2 行起数据，与 config.yaml 的行号配置一致。"""
    path = tmp_path / "input.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows([["说明", "行"] + [""] * (len(HEADERS) - 2), HEADERS, ROW])
    return path


@pytest.fixture
def config_file(tmp_path):
    """基于仓库 config.yaml，指向真实模板，Sheet 检测改为 first。"""
    raw = yaml.safe_load((PROJECT_ROOT / "config.yaml").read_text(encoding="utf-8"))
    raw["sheet_detection"] = {"mode": "first", "pattern": ""}
    for section in (raw["word"], raw["sales"]["word"]):
        section["template_path"] = str(PROJECT_ROOT / section["template_path"])
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(raw, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return path


@pytest.fixture(autouse=True)
def non_interactive(monkeypatch):
    """跳过窗口输入分支，并禁止任何测试弹出文件夹选择窗口。

    Windows 上 tkinter 可用，若某测试选了"2 = 自定义文件夹"而没自己 mock，
    filedialog 会弹出模态窗口等待输入，CI 里无人应答 → 永久阻塞。
    这里默认把它关掉；需要测弹窗的测试自行 monkeypatch 覆盖。
    """
    import src.main as m
    monkeypatch.setattr(sys.stdin, "isatty", lambda: False)
    monkeypatch.setattr(m, "_pick_folder_dialog", lambda initial: (None, False))


def test_output_flag_overrides_directory(excel_file, config_file, tmp_path):
    """-o 指定输出文件夹：两份合同都存到该文件夹，且不改动配置里的默认目录。"""
    custom_dir = tmp_path / "custom_out"

    exit_code = main([str(excel_file), "--config", str(config_file), "-o", str(custom_dir)])

    assert exit_code == 0
    produced = sorted(p.name for p in custom_dir.iterdir())
    assert produced == ["采购合同-供应商B-P01.docx", "销售订单-客户A-S01.docx"]


def test_without_output_flag_uses_config_directory(excel_file, config_file, tmp_path, monkeypatch):
    """不传 -o 时用 config.yaml 的 output.directory。"""
    default_dir = tmp_path / "default_out"
    raw = yaml.safe_load(config_file.read_text(encoding="utf-8"))
    raw["output"]["directory"] = str(default_dir)
    raw["sales"]["output"]["directory"] = str(default_dir)
    config_file.write_text(yaml.safe_dump(raw, allow_unicode=True, sort_keys=False), encoding="utf-8")

    exit_code = main([str(excel_file), "--config", str(config_file)])

    assert exit_code == 0
    assert sorted(p.name for p in default_dir.iterdir()) == [
        "采购合同-供应商B-P01.docx",
        "销售订单-客户A-S01.docx",
    ]


def test_financing_days_from_config_lands_in_sales_contract(excel_file, config_file, tmp_path):
    """非交互时融资时长取 config 默认值 90，写入销售合同。"""
    from docx import Document

    out_dir = tmp_path / "out_days"
    exit_code = main([
        str(excel_file), "--config", str(config_file), "-o", str(out_dir),
    ])

    assert exit_code == 0
    sales_doc = Document(str(out_dir / "销售订单-客户A-S01.docx"))
    text = "\n".join(p.text for p in sales_doc.paragraphs)
    assert "【90】" in text
    assert "{{融资时长}}" not in text


def _interactive(monkeypatch, answers, program_dir):
    """模拟交互式运行：isatty=True、按序返回 answers、程序目录指向临时路径。"""
    import src.main as m
    monkeypatch.setattr(sys.stdin, "isatty", lambda: True)
    monkeypatch.setattr("builtins.input", lambda prompt="": answers.pop(0))
    monkeypatch.setattr(m, "_program_dir", lambda: str(program_dir))


def test_interactive_enter_saves_to_program_dir(excel_file, config_file, tmp_path, monkeypatch):
    """菜单直接回车（=1）：输出到程序同目录。"""
    program_dir = tmp_path / "prog"
    program_dir.mkdir()
    _interactive(monkeypatch, ["", "", "", "", ""], program_dir)

    assert main([str(excel_file), "--config", str(config_file)]) == 0
    assert sorted(p.name for p in program_dir.iterdir()) == [
        "采购合同-供应商B-P01.docx",
        "销售订单-客户A-S01.docx",
    ]


def test_interactive_choice_2_uses_custom_folder(excel_file, config_file, tmp_path, monkeypatch):
    """菜单选 2：输出到自定义文件夹，程序同目录不被写入。"""
    program_dir = tmp_path / "prog"
    program_dir.mkdir()
    custom = tmp_path / "自定义"
    _interactive(monkeypatch, ["", "", "", "", "2", str(custom)], program_dir)

    assert main([str(excel_file), "--config", str(config_file)]) == 0
    assert sorted(p.name for p in custom.iterdir()) == [
        "采购合同-供应商B-P01.docx",
        "销售订单-客户A-S01.docx",
    ]
    assert list(program_dir.iterdir()) == []

def test_interactive_choice_2_uses_folder_dialog(excel_file, config_file, tmp_path, monkeypatch):
    """选 2 时弹出文件夹选择窗口，选中的目录生效。"""
    import src.main as m

    program_dir = tmp_path / "prog"
    program_dir.mkdir()
    picked = tmp_path / "弹窗选的文件夹"
    _interactive(monkeypatch, ["", "", "", "", "2"], program_dir)
    monkeypatch.setattr(m, "_pick_folder_dialog", lambda initial: (str(picked), True))

    assert main([str(excel_file), "--config", str(config_file)]) == 0
    assert sorted(p.name for p in picked.iterdir()) == [
        "采购合同-供应商B-P01.docx",
        "销售订单-客户A-S01.docx",
    ]
    assert list(program_dir.iterdir()) == []


def test_folder_dialog_cancel_falls_back_to_program_dir(excel_file, config_file, tmp_path, monkeypatch):
    """弹窗里点取消：回落到程序同目录，不报错。"""
    import src.main as m

    program_dir = tmp_path / "prog"
    program_dir.mkdir()
    _interactive(monkeypatch, ["", "", "", "", "2"], program_dir)
    monkeypatch.setattr(m, "_pick_folder_dialog", lambda initial: (None, True))

    assert main([str(excel_file), "--config", str(config_file)]) == 0
    assert sorted(p.name for p in program_dir.iterdir()) == [
        "采购合同-供应商B-P01.docx",
        "销售订单-客户A-S01.docx",
    ]


def test_folder_dialog_unavailable_falls_back_to_manual_input(excel_file, config_file, tmp_path, monkeypatch):
    """没有图形环境时不弹窗，改为手动输入路径。"""
    import src.main as m

    program_dir = tmp_path / "prog"
    program_dir.mkdir()
    manual = tmp_path / "手输的文件夹"
    _interactive(monkeypatch, ["", "", "", "", "2", str(manual)], program_dir)
    monkeypatch.setattr(m, "_pick_folder_dialog", lambda initial: (None, False))

    assert main([str(excel_file), "--config", str(config_file)]) == 0
    assert sorted(p.name for p in manual.iterdir()) == [
        "采购合同-供应商B-P01.docx",
        "销售订单-客户A-S01.docx",
    ]
    assert list(program_dir.iterdir()) == []


def test_entered_financing_days_flows_into_premium(excel_file, config_file, tmp_path, monkeypatch):
    """窗口录入的融资时长（120 天）同时写进 {{融资时长}} 并用于溢价计算。"""
    import re
    from docx import Document

    program_dir = tmp_path / "prog"
    program_dir.mkdir()
    # 税点 / 预付款 / 融资时长=120 / 模板=2(溢价版) / 费率=5 / 输出位置=程序同目录
    _interactive(monkeypatch, ["", "", "120", "2", "5", ""], program_dir)

    assert main([str(excel_file), "--config", str(config_file)]) == 0

    produced = [p for p in program_dir.iterdir() if p.name.startswith("销售订单溢价版-")]
    assert len(produced) == 1
    doc = Document(str(produced[0]))
    text = "\n".join(p.text for p in doc.paragraphs)

    # 提货期限用录入的 120 天
    assert "【120】天" in text
    # 溢价 = (采购总额 − 预付款) × 5% ÷ 360 × 120
    # fixture 采购总额 100000、预付款 30% → 垫款 70000
    # 70000 × 0.05 ÷ 360 × 120 = 1166.67
    m = re.search(r"溢价金额【([\d,]+\.\d{2})】", text)
    assert m, "文档里没找到溢价金额"
    assert m.group(1) == "1,166.67"
